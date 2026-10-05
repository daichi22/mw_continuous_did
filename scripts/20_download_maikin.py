"""
20_download_maikin.py
毎月勤労統計調査 地方調査の月次Excelファイルを厚生労働省サイトから一括ダウンロードし、
都道府県×月次パネルデータに変換して data/maikin_panel.csv に保存する。

対象: 事業所規模5人以上、調査産業計（シートT1）
期間: 2013年1月～最新月
ソース: https://www.mhlw.go.jp/toukei/itiran/roudou/monthly/

カラム:
  pref_code   都道府県コード (01-47)
  pref_name   都道府県名
  year        西暦年
  month       月
  workers     常用労働者数（百人/千人）
  hours_total 総実労働時間（時間）
  hours_scheduled 所定内労働時間（時間）
  hours_overtime  所定外労働時間（時間）
  days_worked     出勤日数（日）
  wage_total      現金給与総額（円）
  wage_regular    きまって支給する給与（円）
  wage_scheduled  所定内給与（円）
  wage_special    特別に支払われた給与（円）
  wage_overtime   超過労働給与（円）= wage_regular - wage_scheduled
"""

import time
import io
import sys
import requests
import pandas as pd
from pathlib import Path

SCRIPT_DIR = Path(__file__).parent
PROJECT_ROOT = SCRIPT_DIR.parent
RAW_DIR = PROJECT_ROOT / "data_raw" / "maikin"
RAW_DIR.mkdir(parents=True, exist_ok=True)
OUT_PATH = PROJECT_ROOT / "data" / "maikin_panel.csv"

BASE_URL = "https://www.mhlw.go.jp/toukei/itiran/roudou/monthly"

PREF_NAMES = [
    "北海道", "青森", "岩手", "宮城", "秋田", "山形", "福島",
    "茨城", "栃木", "群馬", "埼玉", "千葉", "東京", "神奈川",
    "新潟", "富山", "石川", "福井", "山梨", "長野",
    "岐阜", "静岡", "愛知", "三重",
    "滋賀", "京都", "大阪", "兵庫", "奈良", "和歌山",
    "鳥取", "島根", "岡山", "広島", "山口",
    "徳島", "香川", "愛媛", "高知",
    "福岡", "佐賀", "長崎", "熊本", "大分", "宮崎", "鹿児島", "沖縄",
]
assert len(PREF_NAMES) == 47


def build_url(year: int, month: int) -> tuple[str, str] | tuple[None, None]:
    """年月からURLとファイル拡張子を返す"""
    mm = f"{month:02d}"

    if year == 2019 and month >= 5:
        # 2019年5-12月: 地方調査ファイルが存在しない期間
        return None, None
    elif year == 2019 and month <= 4:
        # 2019年1-4月: 平成31年
        era_yy = "31"
        return f"{BASE_URL}/{era_yy}/{era_yy}{mm}/xls/{era_yy}{mm}.xls", "xls"
    elif year == 2020 and month >= 4:
        # 2020年4-12月: xls/ ディレクトリ
        reiwa_yy = "02"
        return f"{BASE_URL}/r{reiwa_yy}/{reiwa_yy}{mm}/xls/{reiwa_yy}{mm}.xlsx", "xlsx"
    elif year >= 2020:
        # 2020年1-3月, 2021年以降: xlsx/ ディレクトリ
        reiwa_year = year - 2018
        reiwa_yy = f"{reiwa_year:02d}"
        return f"{BASE_URL}/r{reiwa_yy}/{reiwa_yy}{mm}/xlsx/{reiwa_yy}{mm}.xlsx", "xlsx"
    else:
        # 2012-2018: 平成era
        era_year = year - 1988
        era_yy = f"{era_year:02d}" if era_year < 100 else str(era_year)
        return f"{BASE_URL}/{era_yy}/{era_yy}{mm}/xls/{era_yy}{mm}.xls", "xls"


def download_file(year: int, month: int) -> Path | None:
    """Excelファイルをダウンロードしてローカルに保存"""
    url, ext = build_url(year, month)
    if url is None:
        return None
    local_path = RAW_DIR / f"maikin_{year}{month:02d}.{ext}"

    if local_path.exists() and local_path.stat().st_size > 1000:
        return local_path

    try:
        r = requests.get(url, timeout=30)
        if r.status_code == 200 and len(r.content) > 1000:
            local_path.write_bytes(r.content)
            return local_path
        else:
            print(f"  {year}-{month:02d}: HTTP {r.status_code}, skip")
            return None
    except Exception as e:
        print(f"  {year}-{month:02d}: ERROR {e}")
        return None


def parse_excel(path: Path, year: int, month: int) -> pd.DataFrame | None:
    """T1シート（5人以上、調査産業計）をパースして都道府県データを抽出"""
    ext = path.suffix.lower()

    if ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(str(path))
        sheet_names = wb.sheet_names()
        t1_name = [s for s in sheet_names if "T1" in s]
        if not t1_name:
            print(f"  {year}-{month:02d}: T1 sheet not found in {sheet_names}")
            return None
        ws = wb.sheet_by_name(t1_name[0])

        rows_data = []
        for r in range(ws.nrows):
            row_vals = [ws.cell_value(r, c) for c in range(ws.ncols)]
            rows_data.append(row_vals)
    else:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=True)
        sheet_names = wb.sheetnames
        t1_name = [s for s in sheet_names if "T1" in s]
        if not t1_name:
            print(f"  {year}-{month:02d}: T1 sheet not found in {sheet_names}")
            return None
        ws = wb[t1_name[0]]

        rows_data = []
        for row in ws.iter_rows(values_only=True):
            rows_data.append(list(row))

    pref_rows = []
    seen_prefs = set()
    for row in rows_data:
        if not row or row[0] is None:
            continue
        cell0 = str(row[0]).strip().replace("　", "").replace(" ", "")
        if not cell0 or cell0 in ("全国", "全国計"):
            continue

        has_numeric = any(
            isinstance(v, (int, float)) and v != 0
            for v in row[1:4] if v is not None
        )
        if not has_numeric:
            continue

        matched_pref = None
        for pname in PREF_NAMES:
            clean_pname = pname.replace("　", "").replace(" ", "")
            if cell0 == clean_pname or clean_pname == cell0:
                matched_pref = pname
                break
        if matched_pref is None:
            for pname in PREF_NAMES:
                clean_pname = pname.replace("　", "").replace(" ", "")
                if clean_pname in cell0:
                    matched_pref = pname
                    break

        if matched_pref is None or matched_pref in seen_prefs:
            continue
        seen_prefs.add(matched_pref)

        numeric_vals = []
        for v in row[1:]:
            if v is None or v == "" or str(v).strip() in ("-", "…", "x", "***"):
                numeric_vals.append(None)
            else:
                try:
                    numeric_vals.append(float(v))
                except (ValueError, TypeError):
                    numeric_vals.append(None)

        pref_code = PREF_NAMES.index(matched_pref) + 1
        rec = {
            "pref_code": f"{pref_code:02d}",
            "pref_name": matched_pref,
            "year": year,
            "month": month,
        }

        col_names = [
            "workers", "hours_total", "hours_scheduled",
            "hours_overtime", "days_worked", "wage_total",
            "wage_regular", "wage_scheduled", "wage_special",
        ]
        for i, cname in enumerate(col_names):
            rec[cname] = numeric_vals[i] if i < len(numeric_vals) else None

        if rec["wage_regular"] is not None and rec["wage_scheduled"] is not None:
            rec["wage_overtime"] = rec["wage_regular"] - rec["wage_scheduled"]
        else:
            rec["wage_overtime"] = None

        pref_rows.append(rec)

    if len(pref_rows) != 47:
        print(f"  {year}-{month:02d}: WARNING got {len(pref_rows)} prefs (expected 47)")

    return pd.DataFrame(pref_rows) if pref_rows else None


def main():
    start_year, start_month = 2013, 1
    end_year, end_month = 2026, 12

    all_dfs = []
    fail_count = 0

    for year in range(start_year, end_year + 1):
        for month in range(1, 13):
            if year == start_year and month < start_month:
                continue
            if year == end_year and month > end_month:
                break

            sys.stdout.write(f"\r{year}-{month:02d}...")
            sys.stdout.flush()

            url, _ = build_url(year, month)
            if url is None:
                continue

            path = download_file(year, month)
            if path is None:
                fail_count += 1
                if fail_count >= 3:
                    print(f"\n  3 consecutive failures, stopping at {year}-{month:02d}")
                    break
                continue
            fail_count = 0

            df = parse_excel(path, year, month)
            if df is not None:
                all_dfs.append(df)

            time.sleep(0.3)
        else:
            continue
        break

    if not all_dfs:
        print("\nNo data collected!")
        return

    panel = pd.concat(all_dfs, ignore_index=True)
    panel.sort_values(["pref_code", "year", "month"], inplace=True)
    panel.reset_index(drop=True, inplace=True)

    panel.to_csv(OUT_PATH, index=False, encoding="utf-8-sig")

    n_prefs = panel["pref_code"].nunique()
    n_months = panel.groupby(["year", "month"]).ngroups
    date_range = f"{panel['year'].min()}-{panel['month'].min():02d} to {panel['year'].max()}-{panel.loc[panel['year'] == panel['year'].max(), 'month'].max():02d}"

    print(f"\n\nDone: {len(panel)} rows, {n_prefs} prefectures, {n_months} months")
    print(f"Period: {date_range}")
    print(f"Saved: {OUT_PATH}")
    print(f"\nColumns: {list(panel.columns)}")
    print(f"\nSample:\n{panel.head(10)}")
    print(f"\nMissing values:\n{panel.isnull().sum()}")


if __name__ == "__main__":
    main()
